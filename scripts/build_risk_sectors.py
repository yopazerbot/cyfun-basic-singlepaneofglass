"""Build app/cyfun/framework/risk_model.json from the CCB 'CyFun-Selection' risk assessment workbook.

Usage:
    python scripts/build_risk_sectors.py path/to/BE-NIS2-RA-v20240108.xlsx

The workbook 'Choosing the right Cyber Fundamentals assurance level for your organization'
is published by the Centre for Cybersecurity Belgium (CCB). It is not redistributed here.
This script extracts the scoring criteria and the default threat matrix per NIS2 sector
so the application can reproduce the same calculation.
"""

from __future__ import annotations

import json
import re
import sys
import warnings
from datetime import datetime
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

SKIP = {"Introduction", "Criteria", "Sectors"}
ACTOR_COLS = (5, 7, 9, 11, 13)  # one column per threat actor on each sector sheet


def clean(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def build(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    crit = wb["Criteria"]
    version = crit["E3"].value
    if isinstance(version, datetime):
        version = version.date().isoformat()

    probability = {clean(crit.cell(r, 3).value): float(crit.cell(r, 4).value) for r in (7, 8, 9)}
    prob_text = {clean(crit.cell(r, 6).value).title(): clean(crit.cell(r, 7).value) for r in (7, 8, 9)}
    impact = {clean(crit.cell(r, 3).value): float(crit.cell(r, 4).value) for r in (12, 13, 14)}
    levels = [{"min": float(crit.cell(r, 3).value), "max": float(crit.cell(r, 4).value), "level": clean(crit.cell(r, 5).value)} for r in (17, 18, 19)]
    attack_types = {clean(crit.cell(r, 2).value): int(crit.cell(r, 3).value) for r in (23, 24)}

    sectors = []
    for ws in wb.worksheets:
        if ws.title in SKIP:
            continue
        rows = [
            {
                "attack": clean(ws.cell(r, 2).value),
                "attack_type": int(ws.cell(r, 3).value),
                "impact": clean(ws.cell(r, 4).value),
                "probability": [clean(ws.cell(r, c).value) for c in ACTOR_COLS],
            }
            for r in range(9, 14)
        ]
        sectors.append(
            {
                "id": ws.title,
                "name": clean(ws.cell(6, 2).value),
                "default_size": int(ws.cell(7, 3).value),
                "actors": [clean(ws.cell(7, c).value) for c in ACTOR_COLS],
                "actor_skills": [clean(ws.cell(6, c).value) for c in ACTOR_COLS],
                "rows": rows,
                "workbook_total": float(ws["O14"].value),
                "workbook_level": clean(ws["P14"].value),
            }
        )

    sec = wb["Sectors"]
    sector_list = []
    for r in range(8, sec.max_row + 1):
        code, name, sub = (clean(sec.cell(r, c).value) for c in (1, 2, 3))
        if code and name:
            sector_list.append({"code": code.rstrip("."), "name": name, "subsectors": [sub] if sub else []})
        elif sub and sector_list:
            sector_list[-1]["subsectors"].append(sub)

    return {
        "model": "CyFun-Selection (NIS2 assurance level risk assessment)",
        "owner": "Centre for Cybersecurity Belgium (CCB)",
        "source": {"workbook": path.name, "version": version, "website": "https://www.cyfun.eu"},
        "criteria": {
            "probability": probability,
            "probability_text": prob_text,
            "impact": impact,
            "attack_type": attack_types,
            "organisation_size": {"Large": 3, "Medium": 2, "Small": 1},
            "levels": levels,
            "formula": "score = sum(probability * impact * attack_type * organisation_size) over all attack/actor cells",
        },
        "nis2_sectors": sector_list,
        "sectors": sectors,
    }


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    src = Path(sys.argv[1])
    out = Path(__file__).resolve().parents[1] / "app" / "cyfun" / "framework" / "risk_model.json"
    data = build(src)
    out.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {out}: {len(data['sectors'])} sector sheets, {len(data['nis2_sectors'])} NIS2 sectors")


if __name__ == "__main__":
    main()
