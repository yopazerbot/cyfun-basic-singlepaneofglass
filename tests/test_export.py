"""Workbook export for the three CCB tools. Synthetic workbooks with the CCB layout are built
for CI; when the real CCB workbooks are available next to the repository, the tests also run
against them and verify that nothing but the expected parts changed."""

from __future__ import annotations

import io
import re
import warnings
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from cyfun.export_xlsx import ExportError, ExportInput, col_to_num, fill_workbook, num_to_col
from cyfun.framework import LEVELS, load_framework
from cyfun.scoring import ReqInput, compute

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
ROOT = Path(__file__).resolve().parents[1]
REAL = {
    "BASIC": ROOT.parent / "CyFun2025_ Self-Assessment_tool_BASIC_v2026_02_20.xlsx",
    "IMPORTANT": ROOT.parent / "CyFun2025_ Self-Assessment_tool_IMPORTANT_v2026_02_20.xlsx",
    "ESSENTIAL": ROOT.parent / "CyFun2025_Self-Assessment_tool_ESSENTIAL_v3.1.xlsx",
}
KM_ROW = 40  # synthetic key-measure block start


def shift(col: str, n: int) -> str:
    return num_to_col(col_to_num(col) + n)


def synthetic_template(level: str) -> bytes:
    """Same sheets, rows and formulas as the CCB tool of that level, without styling or protection."""
    fw = load_framework(level)
    lay = fw.layout
    doc, impl = lay["doc_col"], lay["impl_col"]
    sub_doc, sub_impl, cat_doc, cat_impl = shift(doc, 2), shift(doc, 3), shift(doc, 4), shift(doc, 5)
    na = fw.thresholds["na_value"]
    wb = openpyxl.Workbook()
    intro = wb.active
    intro.title = lay["intro_sheet"]
    intro["Q6"] = date(2026, 2, 20)
    intro[lay["date_cell"]] = date(2026, 2, 20)
    sheets = {}
    for name in lay["function_sheets"]:
        ws = wb.create_sheet(name)
        ws["A2"] = "Category"
        ws[f"{lay['req_col']}2"] = "Requirement"
        if lay.get("level_col"):
            ws[f"{lay['level_col']}2"] = "Assurance level"
        sheets[name] = ws
    for r in fw.requirements:
        ws = sheets[r.sheet]
        ws[f"{lay['req_col']}{r.row}"] = f"{r.workbook_id}: {r.text}"
        ws[f"{doc}{r.row}"] = 1
        ws[f"{impl}{r.row}"] = 1
        if lay.get("level_col"):
            ws[f"{lay['level_col']}{r.row}"] = r.level
    for f in fw.functions:
        ws = sheets[f.name]
        for c in f.categories:
            cat_first = c.subcategories[0].requirements[0].row
            # one aggregation cell per group, per dimension, exactly like the CCB formulas (quirks included)
            for dim, col, cat_col, src in (("doc", sub_doc, cat_doc, doc), ("impl", sub_impl, cat_impl, impl)):
                heads = []
                for group in c.groups(dim):
                    rows = [fw.by_id[i].row for i in group]
                    first = rows[0]
                    heads.append(first)
                    parts = [f'IF(OR(${doc}{x}="N/A",${impl}{x}="N/A"),{na:g},${src}{x})' for x in rows]
                    ws[f"{col}{first}"] = f"=AVERAGE({','.join(parts)})" if len(parts) > 1 else f"={parts[0]}"
                ws[f"{cat_col}{cat_first}"] = "=AVERAGE(" + ",".join(f"{col}{x}" for x in heads) + ")"
    summ = wb.create_sheet(lay["summary_sheet"])
    row = 5
    for f in fw.functions:
        for c in f.categories:
            first = c.subcategories[0].requirements[0].row
            summ[f"B{row}"] = c.name
            summ[f"E{row}"] = f"={f.name}!{cat_doc}{first}"
            summ[f"F{row}"] = f"={f.name}!{cat_impl}{first}"
            summ[f"D{row}"] = f"=AVERAGE(E{row},F{row})"
            row += 1
    summ["M4"] = f"=SUM(D5:D{row - 1})/COUNT(D5:D{row - 1})"
    summ["N11"] = f"={lay['intro_sheet']}!{lay['date_cell']}"
    for i, k in enumerate(fw.key_measures):
        r = KM_ROW + i
        summ[f"L{r}"] = k.id
        summ[f"P{r}"] = f"={k.sheet}!{doc}{k.row}"
        summ[f"Q{r}"] = f"={k.sheet}!{impl}{k.row}"
        summ[f"O{r}"] = f"=AVERAGE(P{r},Q{r})"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sample_inputs(fw):
    na_req = next(r for r in fw.requirements if not r.key_measure and not r.management_aspect)
    exp, sco = {}, {}
    for i, r in enumerate(fw.requirements):
        if r.id == na_req.id:
            exp[r.id] = ExportInput(None, None, True, "Not applicable in our context.")
            sco[r.id] = ReqInput(None, None, True)
        else:
            d, m = (i % 5) + 1, ((i * 2) % 5) + 1
            exp[r.id] = ExportInput(d, m, False, f"Justification {r.id} <&> ok")
            sco[r.id] = ReqInput(d, m)
    return exp, sco, na_req


def _check(level: str, template: bytes, synthetic: bool):
    fw = load_framework(level)
    lay = fw.layout
    exp, sco, na_req = sample_inputs(fw)
    out, info = fill_workbook(template, fw, exp, date(2026, 10, 5))
    assert info["requirements_written"] == len(fw.requirements)
    summary = compute(fw, sco)
    wb = openpyxl.load_workbook(io.BytesIO(out), data_only=True)
    ws = wb[lay["summary_sheet"]]
    assert abs(ws["M4"].value - summary.total_maturity) < 1e-9
    assert wb[lay["intro_sheet"]][lay["date_cell"]].value.date() == date(2026, 10, 5)
    for i, c in enumerate(fw.categories):
        g = summary.categories[c.id]
        assert abs(ws[f"E{5 + i}"].value - g.doc) < 1e-9, (level, c.id)
        assert abs(ws[f"F{5 + i}"].value - g.impl) < 1e-9, (level, c.id)
        assert abs(ws[f"D{5 + i}"].value - g.maturity) < 1e-9, (level, c.id)
    # key-measure block: synthetic lists all; the CCB workbooks list a subset in column L
    km = {k.requirement.id: k.maturity for k in summary.key_measures}
    listed = 0
    for r in range(20, ws.max_row + 1):
        rid = str(ws[f"L{r}"].value or "").strip()
        if rid in km:
            assert abs(ws[f"O{r}"].value - km[rid]) < 1e-9, (level, rid)
            listed += 1
    assert listed == (len(km) if synthetic else 13)
    first = next(r for r in fw.requirements if r.id != na_req.id)
    sheet = wb[first.sheet]
    assert sheet[f"{lay['doc_col']}{first.row}"].value == exp[first.id].doc
    assert sheet[f"{lay['comment_col']}{first.row}"].value == f"Justification {first.id} <&> ok"
    na_sheet = wb[na_req.sheet]
    assert na_sheet[f"{lay['doc_col']}{na_req.row}"].value == "N/A" and na_sheet[f"{lay['impl_col']}{na_req.row}"].value == "N/A"
    wbf = openpyxl.load_workbook(io.BytesIO(out))
    assert str(wbf[lay["summary_sheet"]]["M4"].value).startswith("=")
    return out


@pytest.mark.parametrize("level", LEVELS)
def test_fill_synthetic_workbook(level):
    _check(level, synthetic_template(level), synthetic=True)


def test_rejects_wrong_workbook():
    fw = load_framework("BASIC")
    wb = openpyxl.Workbook()
    wb.active.title = "Something"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ExportError):
        fill_workbook(buf.getvalue(), fw, {}, None)
    with pytest.raises(ExportError):
        fill_workbook(b"not a zip", fw, {}, None)


def test_rejects_workbook_of_another_level():
    with pytest.raises(ExportError, match="IMPORTANT"):
        fill_workbook(synthetic_template("BASIC"), load_framework("IMPORTANT"), {}, None)


def test_leaves_unscored_rows_alone():
    fw = load_framework("BASIC")
    out, info = fill_workbook(synthetic_template("BASIC"), fw, {"GV.OC-03.1": ExportInput(4, 3, False, "x")}, None)
    assert info["requirements_written"] == 1
    wb = openpyxl.load_workbook(io.BytesIO(out), data_only=True)
    assert wb["GOVERN"]["F3"].value == 4
    assert wb["GOVERN"]["F4"].value == 1  # template default untouched


@pytest.mark.parametrize("level", LEVELS)
def test_fill_real_ccb_workbook(level):
    path = REAL[level]
    if not path.exists():
        pytest.skip(f"official CCB {level} workbook not available")
    template = path.read_bytes()
    out = _check(level, template, synthetic=False)
    z1, z2 = zipfile.ZipFile(io.BytesIO(template)), zipfile.ZipFile(io.BytesIO(out))
    assert z1.namelist() == z2.namelist()
    changed = {n for n in z1.namelist() if z1.read(n) != z2.read(n)}
    assert all(re.fullmatch(r"xl/workbook\.xml|xl/worksheets/sheet\d+\.xml", n) for n in changed), changed
    assert any("chart" in n for n in z2.namelist())
    assert load_framework(level).layout["summary_sheet"] in openpyxl.load_workbook(io.BytesIO(out), read_only=True).sheetnames
